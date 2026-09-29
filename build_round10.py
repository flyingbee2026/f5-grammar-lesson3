"""Round 10b: two-column Part 2 word table (4 cols x 11 rows) + Part 4 wording fix.
Input: teacher_r10.docx (Helen's hand-finalised file)  ->  Output: teacher.docx (rebuilt)

Layout: fixed 4-column table, widths 2150/2525/2150/2525 twips (= 9350 text width),
word numbering written as literal text (1-10 left, 11-20 right) so the two columns keep
one continuous numbering and no Word list counter is involved.
"""
import copy
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src = docx.Document('teacher_r10.docx')

# ---------- collect the 20 word/noun pairs ----------
noun_tbl = [t for t in src.tables if len(t.rows) >= 15 and t.rows[0].cells[0].text.strip() == 'Word'][0]
pairs = [(noun_tbl.rows[i].cells[0].text.strip(), noun_tbl.rows[i].cells[1].text.strip())
         for i in range(1, len(noun_tbl.rows))]
assert len(pairs) == 20, pairs
assert pairs[0] == ('decide', 'decision') and pairs[-1] == ('strong', 'strength'), pairs

WIDTHS = (2150, 2525, 2150, 2525)      # sums to 9350 = text width


def rpr(bold):
    r = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:cs'), 'Times New Roman')
    r.append(rf)
    if bold:
        r.append(OxmlElement('w:b'))
    return r


def cell(text, width, bold=False):
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(width))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)
    tc.append(tcPr)
    p = OxmlElement('w:p')
    props = rpr(bold)
    pPr = OxmlElement('w:pPr')
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '0')
    sp.set(qn('w:after'), '0')
    sp.set(qn('w:line'), '276')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    pPr.append(props)
    p.append(pPr)
    if text:
        run = OxmlElement('w:r')
        run.append(copy.deepcopy(props))
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        run.append(t)
        p.append(run)
    tc.append(p)
    return tc


tblPr = copy.deepcopy(noun_tbl._tbl.find(qn('w:tblPr')))
tblW = tblPr.find(qn('w:tblW'))
tblW.set(qn('w:w'), '9350')
tblW.set(qn('w:type'), 'dxa')
lay = OxmlElement('w:tblLayout')
lay.set(qn('w:type'), 'fixed')
tblW.addnext(lay)

tbl = OxmlElement('w:tbl')
tbl.append(tblPr)
grid = OxmlElement('w:tblGrid')
for w in WIDTHS:
    gc = OxmlElement('w:gridCol')
    gc.set(qn('w:w'), str(w))
    grid.append(gc)
tbl.append(grid)


def row(cells):
    tr = OxmlElement('w:tr')
    for c in cells:
        tr.append(c)
    return tr


rows = [row([cell('Word', WIDTHS[0], True), cell('Noun', WIDTHS[1], True),
             cell('Word', WIDTHS[2], True), cell('Noun', WIDTHS[3], True)])]
for i in range(10):
    wl, nl = pairs[i]
    wr, nr = pairs[i + 10]
    rows.append(row([cell(f'{i + 1}. {wl}', WIDTHS[0]), cell(nl, WIDTHS[1]),
                     cell(f'{i + 11}. {wr}', WIDTHS[2]), cell(nr, WIDTHS[3])]))
for r in rows:
    tbl.append(r)

old = noun_tbl._tbl
old.addprevious(tbl)
old.getparent().remove(old)

# ---------- Part 4 wording: 'on page 2' -> 'from Part 2' ----------
hits = 0
for p in src.paragraphs:
    for run in p.runs:
        if 'one noun on page 2' in run.text:
            run.text = run.text.replace('one noun on page 2', 'one noun from Part 2')
            hits += 1
assert hits == 1, hits

src.save('teacher.docx')
print('rebuilt teacher.docx: rows', len(rows), 'rows x', len(WIDTHS), 'cols')
