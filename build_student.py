"""Build the STUDENT copy of F5 Grammar Lesson 3 (Nominalisation) from Helen's
finalised TEACHER docx, so the student file mirrors the teacher file exactly
minus answers."""
import copy, re, sys
import docx
from docx.oxml.ns import qn

SRC = 'teacher.docx'
OUT = 'student.docx'

doc = docx.Document(SRC)
body = doc.element.body
children = list(body.iterchildren())

# ---------- 1. drop the teacher's proofreading-answer block ----------
def find_p(pred):
    for ch in body.iterchildren():
        if ch.tag == qn('w:p'):
            from docx.text.paragraph import Paragraph
            p = Paragraph(ch, doc)
            if pred(p):
                return ch, p
    return None, None

ch, p = find_p(lambda p: p.text.strip() == 'Teacher\u2019s answers')
assert ch is not None, 'answers label not found'
answers_tbl = ch.getnext()
assert answers_tbl.tag == qn('w:tbl'), 'table after answers label missing'
blank_before = ch.getprevious()
answers_tbl.getparent().remove(answers_tbl)
body.remove(ch)
if blank_before is not None and blank_before.tag == qn('w:p'):
    body.remove(blank_before)          # avoid a double blank line

# ---------- 2. blank the Noun columns of the forming-nouns table ----------
noun_tbl = None
for t in doc.tables:
    if len(t.rows) >= 10 and t.rows[0].cells[0].text.strip() == 'Word':
        noun_tbl = t
assert noun_tbl is not None
ncols = len(noun_tbl.columns)
noun_cols = [1, 3] if ncols == 4 else [1]
for r in noun_tbl.rows[1:]:
    for ci in noun_cols:
        cell = r.cells[ci]
        for para in cell.paragraphs:
            for run in list(para.runs):
                run._r.getparent().remove(run._r)

# ---------- 3. practice sets: answer row -> blank writing space ----------
set_tables = [t for t in doc.tables if len(t.rows) == 3
              and t.rows[0].cells[0].text.strip() == 'Simple sentence:']
assert len(set_tables) == 8, len(set_tables)
for t in set_tables:
    row = t.rows[2]
    assert 'Teacher\u2019s Answer:' in row.cells[0].text
    merged = row.cells[0].merge(row.cells[1])
    for para in list(merged.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    first = merged.paragraphs[0]
    for run in list(first.runs):
        run._r.getparent().remove(run._r)
    first._p.getparent().remove(first._p)     # drop the label paragraph
    # rebuild two clean empty paragraphs as writing space
    for _ in range(2):
        newp = copy.deepcopy(first._p)
        merged._tc.append(newp)

# ---------- 4. Part 4: drop sample answers, add writing lines ----------
to_drop = []
for ch in body.iterchildren():
    if ch.tag == qn('w:p'):
        from docx.text.paragraph import Paragraph
        txt = Paragraph(ch, doc).text.strip()
        if txt.startswith('Sample answer') or txt.startswith('Each sample opens'):
            to_drop.append(ch)
assert len(to_drop) == 4, len(to_drop)
for ch in to_drop:
    body.remove(ch)

# insertion point: after the last paragraph of the lavender instructions box
instr_last = None
for ch in body.iterchildren():
    if ch.tag == qn('w:p'):
        from docx.text.paragraph import Paragraph
        if 'Use a relative clause' in Paragraph(ch, doc).text:
            instr_last = ch
assert instr_last is not None
# a blank spacer paragraph already exists right after the instructions box
blank = instr_last.getnext()
if not (blank is not None and blank.tag == qn('w:p')
        and not blank.xpath('.//w:t')):
    blank = copy.deepcopy(instr_last)
    for r in blank.xpath('.//w:r'):
        blank.remove(r)
    instr_last.addnext(blank)

template = copy.deepcopy(blank)
lines = template
anchor = blank
for i in range(7):
    newp = copy.deepcopy(template)
    for r in newp.xpath('.//w:r'):
        newp.remove(r)
    r = newp.makeelement(qn('w:r'), {})
    rpr = newp.makeelement(qn('w:rPr'), {})
    rpr.append(newp.makeelement(qn('w:rFonts'), {qn('w:cs'): 'Times New Roman'}))
    r.append(rpr)
    txt = newp.makeelement(qn('w:t'), {})
    txt.text = '_' * 80
    r.append(txt)
    newp.append(r)
    anchor.addnext(newp)
    anchor = newp

doc.save(OUT)
print('saved', OUT)
